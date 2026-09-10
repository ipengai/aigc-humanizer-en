#!/usr/bin/env python3
"""Offline DOCX alignment smoke test with variable mock Huma output.

Usage:
    python tests/run_docx_alignment_fixture.py INPUT.docx OUTPUT_DIR

No network provider is instantiated. The two generated documents simulate a
Huma response with one paragraph removed and one paragraph added, respectively.
"""

import argparse
import hashlib
import json
import os
import zipfile

# Import python-docx before optional application dependencies when this script
# is executed with the bundled artifact runtime.
from docx import Document

from app.helpers.docx_renderer import _replace_paragraph_range
from app.humanizer.failover import FailoverHumanizer
from app.text_extract import extract_text_from_docx, paragraph_list_to_text


class VariableCardinalityMock:
    def __init__(self, scenario):
        self.scenario = scenario
        self.calls = []
        self.changed = False

    def humanize(self, text, mode=None, paragraphs=None):
        self.calls.append(text)
        values = text.split("\n\n")
        if self.scenario == "fewer":
            for index, value in enumerate(values):
                if value.strip().lower() in {"key observations:", "main observation results:"}:
                    del values[index]
                    self.changed = True
                    break
        elif self.scenario == "extra":
            for index, value in enumerate(values):
                boundary = value.find(". ")
                if len(value.split()) >= 40 and boundary > 0:
                    values[index:index + 1] = [
                        value[:boundary + 1], value[boundary + 2:],
                    ]
                    self.changed = True
                    break
        return "\n\n".join(values)


class RecordingFallback:
    def __init__(self):
        self.calls = []

    def humanize(self, text, mode=None, paragraphs=None):
        self.calls.append(text)
        return text


def _write_docx(source_path, output_path, structured):
    document = Document(source_path)
    body_children = list(document.element.body.iterchildren())
    replacements = []
    for item in structured:
        if not item.get("was_rewritten") or item.get("source_format") != "docx":
            continue
        indexes = item.get("source_body_indexes") or []
        if not indexes and item.get("body_index") is not None:
            indexes = [item["body_index"]]
        if indexes:
            replacements.append((indexes, item["text"]))
    replacements.sort(key=lambda value: min(value[0]), reverse=True)
    for indexes, text in replacements:
        _replace_paragraph_range(body_children, indexes, text)
    document.save(output_path)


def _media_hashes(path):
    with zipfile.ZipFile(path) as archive:
        return {
            name: hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist() if name.startswith("word/media/")
        }


def _table_values(document):
    return [
        [[cell.text for cell in row.cells] for row in table.rows]
        for table in document.tables
    ]


def run_scenario(source_path, output_dir, scenario):
    paragraphs = extract_text_from_docx(source_path)
    primary = VariableCardinalityMock(scenario)
    fallback = RecordingFallback()
    humanizer = FailoverHumanizer(primary, fallback)
    _, structured = humanizer.humanize_structured(
        paragraph_list_to_text(paragraphs),
        mode="median",
        paragraphs=paragraphs,
    )
    if not primary.changed:
        raise AssertionError(f"mock scenario did not modify provider output: {scenario}")
    if fallback.calls:
        raise AssertionError(f"alignment unexpectedly used fallback: {scenario}")

    stem = os.path.splitext(os.path.basename(source_path))[0]
    output_path = os.path.join(output_dir, f"{stem}_alignment_{scenario}.docx")
    _write_docx(source_path, output_path, structured)

    source_doc = Document(source_path)
    output_doc = Document(output_path)
    source_text = [paragraph.text for paragraph in source_doc.paragraphs]
    output_text = [paragraph.text for paragraph in output_doc.paragraphs]
    # Rewriting intentionally strips paragraph-edge whitespace. Compare the
    # visible content while still requiring every paragraph slot to survive.
    visible_source_text = [value.strip() for value in source_text]
    visible_output_text = [value.strip() for value in output_text]
    checks = {
        "visible_text_identical": visible_source_text == visible_output_text,
        "paragraph_count": [len(source_text), len(output_text)],
        "table_count": [len(source_doc.tables), len(output_doc.tables)],
        "tables_identical": _table_values(source_doc) == _table_values(output_doc),
        "inline_shape_count": [len(source_doc.inline_shapes), len(output_doc.inline_shapes)],
        "media_identical": _media_hashes(source_path) == _media_hashes(output_path),
        "mock_request_count": len(primary.calls),
        "fallback_request_count": len(fallback.calls),
    }
    if not all((
        checks["visible_text_identical"],
        checks["tables_identical"],
        checks["media_identical"],
        checks["paragraph_count"][0] == checks["paragraph_count"][1],
        checks["table_count"][0] == checks["table_count"][1],
        checks["inline_shape_count"][0] == checks["inline_shape_count"][1],
    )):
        raise AssertionError(json.dumps(checks, ensure_ascii=False, indent=2))
    return output_path, checks


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input_docx")
    parser.add_argument("output_dir")
    args = parser.parse_args()
    os.makedirs(args.output_dir, exist_ok=True)
    report = {}
    for scenario in ("fewer", "extra"):
        output_path, checks = run_scenario(
            os.path.abspath(args.input_docx), os.path.abspath(args.output_dir), scenario,
        )
        report[scenario] = {"output": output_path, **checks}
    report_path = os.path.join(os.path.abspath(args.output_dir), "alignment_report.json")
    with open(report_path, "w", encoding="utf-8") as handle:
        json.dump(report, handle, ensure_ascii=False, indent=2)
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
