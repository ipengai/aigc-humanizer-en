import unittest
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement

from app.helpers.docx_renderer import (
    _replace_paragraph_range,
    _replace_paragraph_text,
)
from app.humanizer.ai_text_humanizer_mock import AITextHumanizerMock


class DocxOutputIntegrityTests(unittest.TestCase):
    def test_mock_humanizer_never_adds_debug_markers(self):
        humanizer = AITextHumanizerMock()
        source = "First paragraph.\n\nSecond paragraph."

        with patch("time.sleep"):
            success, output = humanizer._call_api(source)

        self.assertTrue(success)
        self.assertEqual(output, source)
        self.assertNotIn("++++++++", output)

    def test_replacing_text_removes_stale_inline_tabs(self):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("Old")
        paragraph.runs[0]._r.append(OxmlElement("w:tab"))
        paragraph.add_run("content")

        _replace_paragraph_text(paragraph._p, "New uninterrupted content")

        self.assertEqual(paragraph.text, "New uninterrupted content")
        self.assertFalse(paragraph._p.xpath(".//w:tab"))

    def test_renderer_rejects_paragraph_count_mismatch(self):
        document = Document()
        document.add_paragraph("First")
        document.add_paragraph("Second")
        body_children = list(document.element.body.iterchildren())

        with self.assertRaisesRegex(ValueError, "paragraph mapping mismatch"):
            _replace_paragraph_range(
                body_children, [0, 1],
                "First replacement\n\nSecond replacement\n\nUnexpected extra",
            )


if __name__ == "__main__":
    unittest.main()
